"""Generate sample PDF and DOCX resume documents and extract them (Day 5).

Generates real PDF and DOCX documents in data/raw/resumes/ and executes
the ResumeExtractor engine to produce cleaned outputs in data/extracted/.
"""

from __future__ import annotations

import sys
from pathlib import Path

# Ensure repository root is on sys.path
repo_root = Path(__file__).resolve().parent.parent
if str(repo_root) not in sys.path:
    sys.path.insert(0, str(repo_root))

import docx
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from parsers import ResumeExtractor


def create_sample_docx(file_path: Path) -> None:
    """Create a realistic DOCX resume containing paragraphs, bullets, and tables."""
    doc = docx.Document()

    # Title & Contact
    title = doc.add_paragraph()
    run = title.add_run("ALEX CHEN")
    run.bold = True
    run.font.size = docx.shared.Pt(18)

    contact = doc.add_paragraph("Email: alex.chen.dev@example.com | Phone: +1-415-555-0192 | Location: San Francisco, CA")
    contact.runs[0].font.size = docx.shared.Pt(10)

    # Professional Summary
    doc.add_heading("PROFESSIONAL SUMMARY", level=1)
    doc.add_paragraph(
        "Senior Backend Engineer with 7+ years of experience designing high-throughput, "
        "fault-tolerant distributed systems. Specialized in Go, Python, gRPC microservices, "
        "and event-driven architectures processing 50M+ events daily."
    )

    # Technical Skills Table
    doc.add_heading("TECHNICAL SKILLS", level=1)
    table = doc.add_table(rows=1, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Category"
    hdr_cells[1].text = "Technologies / Skills"

    skills_data = [
        ("Languages", "Go, Python, SQL, C++, Bash"),
        ("Distributed Systems", "Apache Kafka, gRPC, RabbitMQ, Redis, Elasticsearch"),
        ("Cloud & Infrastructure", "AWS (EKS, RDS, S3), Kubernetes, Docker, Terraform"),
        ("Databases", "PostgreSQL, MongoDB, DynamoDB"),
    ]
    for cat, items in skills_data:
        row_cells = table.add_row().cells
        row_cells[0].text = cat
        row_cells[1].text = items

    # Work Experience
    doc.add_heading("WORK EXPERIENCE", level=1)
    p_exp1 = doc.add_paragraph()
    p_exp1.add_run("Staff Distributed Systems Engineer | CloudScale Inc. (2021 - Present)").bold = True

    bullets = [
        "Architected an event-driven payment reconciliation service in Go processing $40M daily volume.",
        "Reduced API p99 latency from 320ms to 45ms by introducing Redis caching clusters and connection pooling.",
        "Mentored a pod of 5 backend engineers in concurrency patterns, unit testing, and observability.",
    ]
    for bullet in bullets:
        doc.add_paragraph(f"- {bullet}")

    # Education
    doc.add_heading("EDUCATION", level=1)
    doc.add_paragraph("Bachelor of Science in Computer Science | UC Berkeley (2013 - 2017) | GPA: 3.8/4.0")

    doc.save(str(file_path))
    print(f"[OK] Generated DOCX resume: {file_path}")


def create_sample_pdf(file_path: Path) -> None:
    """Create a realistic PDF resume using ReportLab with headings, bullets, and metadata."""
    doc = SimpleDocTemplate(
        str(file_path),
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )
    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        "ResumeTitle",
        parent=styles["Heading1"],
        fontSize=18,
        leading=22,
        spaceAfter=4,
    )
    heading_style = ParagraphStyle(
        "ResumeHeading",
        parent=styles["Heading2"],
        fontSize=12,
        leading=16,
        spaceBefore=10,
        spaceAfter=4,
        textColor="#1e293b",
    )
    body_style = ParagraphStyle(
        "ResumeBody",
        parent=styles["Normal"],
        fontSize=10,
        leading=14,
        spaceAfter=4,
    )
    bullet_style = ParagraphStyle(
        "ResumeBullet",
        parent=styles["Normal"],
        fontSize=9.5,
        leading=13,
        leftIndent=15,
        spaceAfter=3,
    )

    story = [
        Paragraph("<b>PRIYA SHARMA</b>", title_style),
        Paragraph("Email: priya.sharma.tech@example.com | Phone: +1-206-555-0144 | Location: Seattle, WA", body_style),
        Spacer(1, 8),
        Paragraph("<b>PROFESSIONAL SUMMARY</b>", heading_style),
        Paragraph(
            "Lead Frontend Engineer with 6+ years driving design systems, accessible UI/UX, "
            "and high-performance Web applications using React, TypeScript, Next.js, and GraphQL.",
            body_style,
        ),
        Spacer(1, 6),
        Paragraph("<b>TECHNICAL SKILLS</b>", heading_style),
        Paragraph("• <b>Frontend:</b> React, Next.js, TypeScript, Tailwind CSS, Vue.js, Webpack, Vite", bullet_style),
        Paragraph("• <b>Testing & Tools:</b> Playwright, Jest, Storybook, Figma, Git, GraphQL", bullet_style),
        Spacer(1, 6),
        Paragraph("<b>WORK EXPERIENCE</b>", heading_style),
        Paragraph("<b>Lead Frontend Engineer | PixelCraft Media (2022 - Present)</b>", body_style),
        Paragraph("• Led a team of 6 engineers revamping customer dashboard in Next.js 14, improving Core Web Vitals.", bullet_style),
        Paragraph("• Created an enterprise accessible component library reaching WCAG 2.1 AA compliance.", bullet_style),
        Paragraph("• Improved Largest Contentful Paint (LCP) by 65% across primary web portals.", bullet_style),
        Spacer(1, 6),
        Paragraph("<b>EDUCATION</b>", heading_style),
        Paragraph("Bachelor of Engineering in Information Technology | University of Washington (2014 - 2018)", body_style),
    ]

    doc.build(story)
    print(f"[OK] Generated PDF resume: {file_path}")


def main() -> None:
    """Generate documents and execute extraction pipeline."""
    raw_dir = repo_root / "data" / "raw" / "resumes"
    extracted_dir = repo_root / "data" / "extracted"
    raw_dir.mkdir(parents=True, exist_ok=True)
    extracted_dir.mkdir(parents=True, exist_ok=True)

    # 1. Generate DOCX and PDF test resumes
    docx_path = raw_dir / "resume_01_backend.docx"
    pdf_path = raw_dir / "resume_02_frontend.pdf"

    create_sample_docx(docx_path)
    create_sample_pdf(pdf_path)

    # 2. Extract both using ResumeExtractor and store cleaned files
    extractor = ResumeExtractor()

    print("\nRunning ResumeExtractor on generated documents...")
    for doc_path in [docx_path, pdf_path]:
        txt_out, json_out, result = extractor.extract_and_save(doc_path, extracted_dir)
        print(f"  Processed {result.file_type} [{result.file_name}]:")
        print(f"    -> Cleaned Text: {txt_out.name} ({result.word_count} words)")
        print(f"    -> Metadata JSON: {json_out.name} (Sections detected: {result.detected_sections})")

    # Also extract any existing text resumes for completeness
    for txt_file in sorted(raw_dir.glob("*.txt"))[:3]:
        t_out, j_out, res = extractor.extract_and_save(txt_file, extracted_dir)
        print(f"  Processed TXT [{res.file_name}] -> {t_out.name} ({res.word_count} words)")

    print("\nAll sample documents generated and extracted successfully!")


if __name__ == "__main__":
    main()
